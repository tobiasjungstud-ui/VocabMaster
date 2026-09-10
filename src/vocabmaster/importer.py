"""Aus der Excel-Wortliste des Lehrmittels wird die Datenbank.

Der Import läuft einmal und schreibt **eine JSON-Datei je Unit** plus eine
``index.json``. Warum nicht SQLite als Datei: Die Wortliste wird von Hand
nachgepflegt und im Git-Verlauf gelesen; eine Binärdatei zeigt im Diff nur
"binary files differ", eine JSON-Datei je Unit zeigt genau, welches Wort
dazukam. Für Abfragen baut :mod:`vocabmaster.database` daraus beim Öffnen
eine SQLite-Datenbank im Arbeitsspeicher - so bleibt der Git-Verlauf lesbar
und SQL trotzdem verfügbar.

Jede Unit-Datei nennt Unit-Nummer, Thema, Quelldatei und Importdatum, damit
jedes erzeugte Dokument auf eine benannte Fassung der Wortliste zurückgeführt
werden kann.

Alles, was zur Laufzeit teuer wäre, wird hier vorberechnet - vor allem die
Häufigkeitswerte (Zipf), damit die eigentliche Anwendung ohne ``wordfreq``
auskommt und in jeder Umgebung gleich rechnet.

Die Wortlisten der Lehrmittel sind nie sauber tabellarisch. Behandelt werden
deshalb gezielt:

* Titelzeilen oberhalb der Kopfzeile
* zwei Ebenen von Blocküberschriften: erst der Bereich (``Culture``,
  ``Curriculum extra``, ``Project``, ``Literature``, ``Extra listening and
  speaking``), darunter die Unit (``Unit 1`` … ``Unit 8``)
* Sektionsangaben mit Seitenzahl (``Starter Unit, p.4``) - die Zahl am Ende
  ist die Seite, nicht die Unit
* Kurzformen und Sonderfälle (``CE - Unit 5``, ``WB, Unit 4``,
  ``Answer: Unit 4``)
* doppelte Zeilen und Zeilen ohne Übersetzung
"""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import asdict, dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

SCHEMA_VERSION = 1

# --------------------------------------------------------------------------
# Spaltenerkennung
# --------------------------------------------------------------------------
_HEADER_ALIASES = {
    "english": {"english", "englisch", "word", "wort", "englisches wort", "en"},
    "german": {
        "translation", "german", "deutsch", "übersetzung", "uebersetzung",
        "deutsche übersetzung", "de",
    },
    "section": {"section", "sektion", "unit", "abschnitt", "quelle", "fundstelle"},
    "oxford": {"oxford 3000", "oxford3000", "oxford", "oxford 5000"},
    "pos": {"part of speech", "wortart", "pos"},
    "pronunciation": {"pronunciation", "aussprache", "lautschrift"},
    "example": {"example", "beispiel", "beispielsatz", "example sentence"},
}

# --------------------------------------------------------------------------
# Sektionsangaben
# --------------------------------------------------------------------------
_PAGE_SUFFIX = re.compile(
    r"[,;]\s*(?:p{1,2}\.?|pages?|seiten?|s\.)\s*[\dIVXivx]+\s*[-–—]?\s*[\dIVXivx]*\s*$",
    re.IGNORECASE,
)
_PAGE_PAREN = re.compile(r"\(\s*(?:p{1,2}\.?|s\.)\s*[^)]*\)\s*$", re.IGNORECASE)
_PAGE_NUMBER = re.compile(r"p{1,2}\.?\s*(\d{1,3})", re.IGNORECASE)

#: Bereich -> (Kennung, Reihenfolge im Bericht). ``hauptteil`` ist der Block
#: ``Unit N`` selbst; alles andere ist Zusatzmaterial und zählt nach den
#: Arbeitsanweisungen **nicht** zum Umfang einer Unit.
_KINDS: tuple[tuple[str, str], ...] = (
    (r"extra\s+listening(?:\s+and\s+speaking)?", "extra_listening"),
    (r"curriculum\s+extra|^ce\b", "curriculum_extra"),
    (r"culture|kultur", "culture"),
    (r"project", "project"),
    (r"literature|literatur", "literature"),
    (r"song", "songs"),
    (r"workbook|^wb\b", "workbook"),
    (r"answer", "answer_key"),
    (r"starter", "starter"),
    (r"unit|einheit", "hauptteil"),
)

KIND_LABELS = {
    "hauptteil": "Hauptteil",
    "starter": "Starter Unit",
    "culture": "Culture",
    "curriculum_extra": "Curriculum extra",
    "extra_listening": "Extra Listening and Speaking",
    "project": "Project",
    "literature": "Literature",
    "songs": "Songs",
    "workbook": "Workbook",
    "answer_key": "Answer key",
    "unbekannt": "ohne Angabe",
}

_UNIT_NUMBER = re.compile(r"(?:unit|einheit|term)?\s*(\d{1,2})\s*$", re.IGNORECASE)
_ANY_NUMBER = re.compile(r"(\d{1,2})")
_TERM = re.compile(r"\bterm\s*(\d{1,2})\b", re.IGNORECASE)

#: Höchste plausible Unit-Nummer. Alles darüber ist eine Seitenzahl, die
#: der Abtrennung entgangen ist.
MAX_UNIT = 20

_POS_MAP = {
    "n": "noun", "n pl": "noun", "npl": "noun",
    "v": "verb", "phr v": "verb", "phrv": "verb",
    "adj": "adj", "adv": "adv", "conj": "conj", "prep": "prep",
    "pron": "pron", "det": "det", "exclam": "exclam", "number": "number",
}


def strip_page_reference(section: str) -> str:
    """``'Starter Unit, p.4'`` -> ``'Starter Unit'``."""
    s = str(section).replace("\n", " ").strip()
    previous = None
    while previous != s:
        previous = s
        s = _PAGE_SUFFIX.sub("", s).strip()
        s = _PAGE_PAREN.sub("", s).strip()
    return s.strip().strip(",;").strip()


def page_of(section: str) -> int | None:
    m = _PAGE_NUMBER.search(str(section))
    return int(m.group(1)) if m else None


def parse_section(section: str) -> tuple[str, int | None, str]:
    """``'Culture 7'`` -> ``('culture', 7, 'Culture 7')``.

    Gibt ``(kind, unit, label)`` zurück. ``unit`` ist ``0`` für die Starter
    Unit und ``None``, wenn keine Unit erkennbar ist.
    """
    raw = str(section).replace("\n", " ").strip()
    if not raw:
        return "unbekannt", None, ""

    # Mehrfachangaben wie "WB, Unit 4, p.42; WB, Unit 3, p.74;" zuerst
    # zerlegen - es gilt die erste. Erst danach die Seitenangabe abtrennen,
    # sonst bleibt "p.42" stehen und wird als Unit 42 gelesen.
    first = strip_page_reference(re.split(r"\s*;\s*", raw)[0].strip().strip(","))
    if not first:
        return "unbekannt", None, ""
    lowered = first.lower()

    kind = "unbekannt"
    for pattern, name in _KINDS:
        if re.search(pattern, lowered):
            kind = name
            break

    if kind == "starter":
        return "starter", 0, first

    unit: int | None = None
    for pattern in (_TERM, _UNIT_NUMBER, _ANY_NUMBER):
        m = pattern.search(first)
        if m:
            unit = int(m.group(1))
            break
    # Eine zweistellige Zahl über MAX_UNIT ist keine Unit, sondern eine
    # übersehene Seitenzahl. Lieber nicht zuordnen als falsch zuordnen.
    if unit is not None and unit > MAX_UNIT:
        unit = None
    return kind, unit, first


def normalise_pos(raw: str) -> str:
    """``'(phr v)'`` -> ``'verb'``; Mehrfachangaben nehmen die erste."""
    text = re.sub(r"[()]", " ", str(raw or "")).strip().lower()
    if not text:
        return ""
    first = re.split(r"\s*/\s*", text)[0].strip()
    return _POS_MAP.get(first, _POS_MAP.get(first.replace(" ", ""), ""))


# --------------------------------------------------------------------------
# Datensatz
# --------------------------------------------------------------------------
@dataclass
class Row:
    """Eine Zeile der Wortliste, so wie sie in der Datenbank landet."""

    english: str
    german: str
    unit: int | None
    kind: str
    section: str
    page: int | None = None
    pos: str = ""
    pronunciation: str = ""
    example: str = ""
    oxford3000: bool = False
    zipf: float = 0.0
    headword: str = ""
    row: int = 0

    @property
    def is_core(self) -> bool:
        """Steht der Eintrag im Hauptteil der Unit (nicht in Culture, Project …)?"""
        return self.kind in ("hauptteil", "starter")


@dataclass
class ImportResult:
    rows: list[Row] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    source: str = ""
    sheet: str = ""
    checksum: str = ""
    imported: str = ""

    def units(self) -> list[int]:
        return sorted({r.unit for r in self.rows if r.unit is not None})


# --------------------------------------------------------------------------
# Einlesen
# --------------------------------------------------------------------------
def _norm(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"[ \t]+", " ", str(value)).replace("\n", " ").strip()


def read_grid(path: str | Path) -> tuple[list[list[str]], str]:
    """Liest ``.xls`` und ``.xlsx`` in ein Raster aus Zeichenketten."""
    p = Path(path)
    if p.suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(p))
        sheet = book.sheet_by_index(0)
        grid = [
            [_norm(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
            for r in range(sheet.nrows)
        ]
        return grid, sheet.name

    from openpyxl import load_workbook

    wb = load_workbook(str(p), data_only=True, read_only=True)
    try:
        ws = wb.worksheets[0]
        grid = [[_norm(v) for v in row] for row in ws.iter_rows(values_only=True)]
        return grid, ws.title
    finally:
        wb.close()


def _find_header(grid: list[list[str]]) -> tuple[int, dict[str, int]]:
    for index, row in enumerate(grid[:40]):
        mapping: dict[str, int] = {}
        for column, cell in enumerate(row):
            text = cell.lower().strip("*: ").split("\n")[0].strip()
            if not text:
                continue
            for name, aliases in _HEADER_ALIASES.items():
                if text in aliases:
                    mapping.setdefault(name, column)
        if "english" in mapping and "german" in mapping:
            return index, mapping
    raise ValueError(
        "Keine Kopfzeile mit den Spalten 'English' und 'Translation'/'Deutsch' "
        "gefunden. Bitte das Tabellenblatt prüfen."
    )


#: Überschriften, die einen Bereich eröffnen statt eine Unit.
_AREA_HEADING = re.compile(
    r"^\s*(extra\s+listening|curriculum\s+extra|culture|project|literature|song)",
    re.IGNORECASE,
)


def _zipf_lookup():
    """``wordfreq`` wenn vorhanden - sonst ein neutraler Wert mit Warnung."""
    try:
        from wordfreq import zipf_frequency
    except ImportError:  # pragma: no cover - nur beim Import relevant
        return None

    def lookup(word: str) -> float:
        tokens = [t for t in re.findall(r"[a-z'’-]+", word.lower()) if len(t) > 2]
        if not tokens:
            return 3.8
        scores = [zipf_frequency(t, "en") for t in tokens]
        scores = [s for s in scores if s > 0] or [2.0]
        value = min(scores)
        if len(word.split()) > 1:
            value -= 0.60  # Wendungen sind schwerer als ihr häufigstes Wort
        return round(value, 3)

    return lookup


def import_wordlist(path: str | Path) -> ImportResult:
    """Liest die Wortliste und ordnet jeden Eintrag Unit und Bereich zu."""
    from .list.normalize import headword as _headword

    grid, sheet = read_grid(path)
    header_index, cols = _find_header(grid)
    result = ImportResult(
        source=Path(path).name,
        sheet=sheet,
        checksum=file_checksum(path),
        imported=date.today().isoformat(),
    )

    lookup = _zipf_lookup()
    if lookup is None:
        result.warnings.append(
            "wordfreq ist nicht installiert - alle Häufigkeitswerte stehen auf "
            "3.8. Für einen belastbaren Import bitte 'pip install wordfreq'."
        )

    en, de = cols["english"], cols["german"]
    sec = cols.get("section")
    area: str | None = None          # 'Culture', 'Project', …
    block_unit: int | None = None    # aus der Blocküberschrift 'Unit 7'
    block_kind = "hauptteil"
    seen: set[tuple[str, str, int | None, str]] = set()
    mismatches: list[str] = []
    missing_section = 0

    for offset, raw in enumerate(grid[header_index + 1 :], start=header_index + 2):
        row = list(raw) + [""] * (max(cols.values()) + 1 - len(raw))
        english, german = row[en], row[de]

        if not english and not german:
            # Blocküberschrift: Text irgendwo in der Zeile, aber keine Vokabel.
            text = next((c for c in row if c), "")
            if not text:
                continue
            if _AREA_HEADING.match(text):
                area, block_unit, block_kind = text, None, parse_section(text)[0]
            else:
                kind, unit, _ = parse_section(text)
                block_unit = unit
                # 'Unit 7' unter 'Culture' bedeutet Culture 7, nicht Hauptteil.
                block_kind = block_kind if area and kind == "hauptteil" else kind
            continue

        if not english:
            continue
        if not german:
            result.warnings.append(
                f"Zeile {offset}: '{english}' hat keine deutsche Übersetzung "
                "und wird übersprungen."
            )
            continue

        section = row[sec] if sec is not None else ""
        if section:
            kind, unit, label = parse_section(section)
        else:
            missing_section += 1
            kind, unit, label = block_kind, block_unit, (area or "")

        # Die Blocküberschrift ist die zweite, unabhängige Quelle.
        if unit is None and block_unit is not None:
            unit, kind = block_unit, block_kind
        elif unit is not None and block_unit is not None and unit != block_unit:
            mismatches.append(f"Zeile {offset} ('{english}'): Block {block_unit}, Spalte {unit}")

        headword = _headword(english)
        key = (english.lower(), german.lower(), unit, kind)
        if key in seen:
            continue
        seen.add(key)

        result.rows.append(
            Row(
                english=english,
                german=german,
                unit=unit,
                kind=kind,
                section=label or section,
                page=page_of(section),
                pos=normalise_pos(row[cols["pos"]]) if "pos" in cols else "",
                pronunciation=row[cols["pronunciation"]] if "pronunciation" in cols else "",
                example=row[cols["example"]] if "example" in cols else "",
                oxford3000=bool(row[cols["oxford"]].strip()) if "oxford" in cols else False,
                zipf=lookup(headword) if lookup else 3.8,
                headword=headword,
                row=offset,
            )
        )

    if mismatches:
        result.warnings.append(
            f"{len(mismatches)} Einträge: Blocküberschrift und Spalte 'Section' "
            f"nennen unterschiedliche Units. Es gilt 'Section'. "
            f"Beispiele: {'; '.join(mismatches[:3])}"
        )
    if missing_section:
        result.warnings.append(
            f"{missing_section} Einträge ohne Angabe in der Spalte 'Section' - "
            "für sie gilt die Blocküberschrift."
        )
    unassigned = sum(1 for r in result.rows if r.unit is None)
    if unassigned:
        result.warnings.append(
            f"{unassigned} Einträge konnten keiner Unit zugeordnet werden."
        )
    if not result.rows:
        raise ValueError("Die Datei enthält keine auswertbaren Vokabelzeilen.")
    return result


def file_checksum(path: str | Path) -> str:
    """SHA-256 der Quelldatei - macht nachprüfbar, aus welcher Excel-Fassung
    die Datenbank stammt."""
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 16), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_themes(path: str | Path | None = None) -> dict[int, dict[str, Any]]:
    """Thema und Titel je Unit. Die Datei wird gelesen, nie überschrieben."""
    target = Path(path or Path(__file__).resolve().parent / "data" / "themen.json")
    if not target.exists():
        return {}
    payload = json.loads(target.read_text(encoding="utf-8"))
    return {int(k): v for k, v in payload.get("themen", {}).items()}


def unit_filename(unit: int | None) -> str:
    return "unit_ohne.json" if unit is None else f"unit_{unit:02d}.json"


def write_database(
    result: ImportResult, directory: str | Path, themes: dict | None = None
) -> list[Path]:
    """Schreibt je Unit eine JSON-Datei plus ``index.json``.

    Vorhandene Unit-Dateien, die in der neuen Quelle nicht mehr vorkommen,
    werden **gelöscht**: Aus einer veralteten Wortliste darf nichts
    überleben, was in der aktuellen nicht mehr steht.
    """
    themes = themes if themes is not None else load_themes()
    target = Path(directory)
    target.mkdir(parents=True, exist_ok=True)

    by_unit: dict[int | None, list[Row]] = {}
    for row in result.rows:
        by_unit.setdefault(row.unit, []).append(row)

    written: list[Path] = []
    index_units = []
    for unit in sorted(by_unit, key=lambda u: (u is None, u)):
        rows = by_unit[unit]
        theme = themes.get(unit, {}) if unit is not None else {}
        payload = {
            "schema": SCHEMA_VERSION,
            "unit": unit,
            "titel": theme.get("titel") or (f"Unit {unit}" if unit else "ohne Unit"),
            "thema": theme.get("thema", ""),
            "seiten": theme.get("seiten", ""),
            "leitwoerter": theme.get("leitwoerter", []),
            "quelle": {
                "datei": result.source,
                "tabellenblatt": result.sheet,
                "pruefsumme_sha256": result.checksum,
                "importiert": result.imported,
            },
            "anzahl": len(rows),
            "anzahl_hauptteil": sum(1 for r in rows if r.is_core),
            "woerter": [asdict(r) for r in sorted(rows, key=lambda r: r.row)],
        }
        path = target / unit_filename(unit)
        path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=1) + "\n", encoding="utf-8"
        )
        written.append(path)
        index_units.append(
            {
                "unit": unit,
                "datei": path.name,
                "titel": payload["titel"],
                "thema": payload["thema"],
                "anzahl": payload["anzahl"],
                "anzahl_hauptteil": payload["anzahl_hauptteil"],
            }
        )

    index = {
        "schema": SCHEMA_VERSION,
        "quelle": {
            "datei": result.source,
            "tabellenblatt": result.sheet,
            "pruefsumme_sha256": result.checksum,
            "importiert": result.imported,
        },
        "eintraege": len(result.rows),
        "units": index_units,
        "hinweise": result.warnings,
    }
    index_path = target / "index.json"
    index_path.write_text(
        json.dumps(index, ensure_ascii=False, indent=1) + "\n", encoding="utf-8"
    )
    written.append(index_path)

    # Reste einer früheren, anderen Wortliste entfernen (§ "keine Vermischung").
    keep = {p.name for p in written}
    for stale in target.glob("unit_*.json"):
        if stale.name not in keep:
            stale.unlink()
            result.warnings.append(
                f"{stale.name} stammte aus einer früheren Wortliste und wurde gelöscht."
            )
    return written
