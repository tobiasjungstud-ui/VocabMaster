"""Erzeugt die Word-Datei im Layout der Vorlage.

Alle Masse stammen aus der bereitgestellten Vorlagendatei und werden hier
unverändert übernommen: Seitenränder, Tabellenbreite, Spaltenbreiten,
Zeilenhöhe, Rahmenlinien, Schriftart und Schriftgrad.
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from io import BytesIO
from pathlib import Path
from typing import IO

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

from ..config import Settings
from .layout import FitResult, check_fits_page
from .models import TestItem, TestPair

# --- Masse aus der Vorlage (in Twips bzw. DXA) ---------------------------
PAGE_WIDTH = 11906
PAGE_HEIGHT = 16838
MARGIN_TOP = 728
MARGIN_RIGHT = 1417
MARGIN_BOTTOM = 625
MARGIN_LEFT = 1417
HEADER_DISTANCE = 419
FOOTER_DISTANCE = 0

TABLE_WIDTH = 9520
COLUMN_WIDTHS = (435, 2511, 2024, 4550)
CELL_MARGIN = 70
ROW_HEIGHT = 397
BORDER_SIZE = 4
BORDER_COLOR = "000000"

_FALLBACK_FONTS = "Century Gothic, Calibri, sans-serif"


def _set(element, tag: str, **attrs) -> None:
    child = OxmlElement(tag)
    for key, value in attrs.items():
        child.set(qn(f"w:{key}"), str(value))
    element.append(child)


def _configure_section(document: Document) -> None:
    """Seitenformat exakt wie in der Vorlage."""
    section = document.sections[0]
    section.page_width = Twips(PAGE_WIDTH)
    section.page_height = Twips(PAGE_HEIGHT)
    section.top_margin = Twips(MARGIN_TOP)
    section.right_margin = Twips(MARGIN_RIGHT)
    section.bottom_margin = Twips(MARGIN_BOTTOM)
    section.left_margin = Twips(MARGIN_LEFT)
    section.header_distance = Twips(HEADER_DISTANCE)
    section.footer_distance = Twips(FOOTER_DISTANCE)


def _style_run(run, settings: Settings, bold: bool) -> None:
    """Century Gothic, 11 pt, schwarz - wie jede Zelle der Vorlage."""
    run.font.name = settings.font_name
    run.font.size = Pt(settings.font_size_pt)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("ascii", "hAnsi"):
        fonts.set(qn(f"w:{attr}"), settings.font_name)
    fonts.set(qn("w:eastAsia"), "Times New Roman")
    fonts.set(qn("w:cs"), "Times New Roman")
    _set(rpr, "w:color", val=BORDER_COLOR)
    _set(rpr, "w:kern", val="0")


def _prepare_paragraph(paragraph) -> None:
    """Kein Abstand, einfacher Zeilenabstand - wie in der Vorlage."""
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    ppr = paragraph._p.get_or_add_pPr()
    spacing = ppr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        ppr.append(spacing)
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")


def _cell_borders(cell, top: bool, bottom: bool) -> None:
    """Nur waagerechte Linien - senkrechte Rahmen hat die Vorlage nicht."""
    tc_pr = cell._tc.get_or_add_tcPr()
    for existing in tc_pr.findall(qn("w:tcBorders")):
        tc_pr.remove(existing)
    borders = OxmlElement("w:tcBorders")
    for edge, present in (("top", top), ("left", False), ("bottom", bottom), ("right", False)):
        element = OxmlElement(f"w:{edge}")
        if present:
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), str(BORDER_SIZE))
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), BORDER_COLOR)
        else:
            element.set(qn("w:val"), "nil")
        borders.append(element)
    tc_pr.append(borders)


def _write_cell(
    cell,
    settings: Settings,
    *,
    text: str = "",
    bold: bool = False,
    align_right: bool = False,
    highlight: str = "",
    top_border: bool = True,
    bottom_border: bool = True,
    width: int = 0,
) -> None:
    """Füllt eine Zelle; ``highlight`` wird innerhalb des Texts fett gesetzt."""
    if width:
        cell.width = Twips(width)
        tc_pr = cell._tc.get_or_add_tcPr()
        for existing in tc_pr.findall(qn("w:tcW")):
            tc_pr.remove(existing)
        _set(tc_pr, "w:tcW", w=width, type="dxa")

    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _cell_borders(cell, top_border, bottom_border)

    paragraph = cell.paragraphs[0]
    _prepare_paragraph(paragraph)
    if align_right:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for chunk, is_bold in _split_highlight(text, highlight):
        run = paragraph.add_run(chunk)
        _style_run(run, settings, bold or is_bold)


def _split_highlight(text: str, highlight: str) -> list[tuple[str, bool]]:
    """Zerlegt den Satz in ``(teiltext, fett)`` - das Zielwort wird fett."""
    if not highlight or not text:
        return [(text, False)] if text else [("", False)]
    index = text.lower().find(highlight.lower())
    if index < 0:
        # Zielwort nicht wörtlich gefunden: als Ganzes ohne Hervorhebung setzen.
        return [(text, False)]
    parts: list[tuple[str, bool]] = []
    if index:
        parts.append((text[:index], False))
    parts.append((text[index : index + len(highlight)], True))
    tail = text[index + len(highlight) :]
    if tail:
        parts.append((tail, False))
    return parts


def _set_table_properties(table) -> None:
    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblCellMar", "w:tblLook", "w:tblBorders"):
        for existing in tbl_pr.findall(qn(tag)):
            tbl_pr.remove(existing)

    _set(tbl_pr, "w:tblW", w=TABLE_WIDTH, type="dxa")

    margins = OxmlElement("w:tblCellMar")
    for edge in ("left", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(CELL_MARGIN))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tbl_pr.append(margins)

    look = OxmlElement("w:tblLook")
    for key, value in (
        ("val", "04A0"), ("firstRow", "1"), ("lastRow", "0"),
        ("firstColumn", "1"), ("lastColumn", "0"), ("noHBand", "0"), ("noVBand", "1"),
    ):
        look.set(qn(f"w:{key}"), value)
    tbl_pr.append(look)

    # Spaltenraster exakt setzen
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        table._tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for width in COLUMN_WIDTHS:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    table._tbl.insert(1, grid)


def _set_row_height(row, height: int = ROW_HEIGHT) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    for existing in tr_pr.findall(qn("w:trHeight")):
        tr_pr.remove(existing)
    _set(tr_pr, "w:trHeight", val=height)


def _add_heading(document: Document, text: str, page_break: bool = False):
    paragraph = document.add_paragraph()
    _prepare_paragraph(paragraph)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run()
    if page_break:
        run.add_break(WD_BREAK.PAGE)
    run = paragraph.add_run(text)
    run.font.size = Pt(12)
    run.font.bold = True
    return paragraph


def _add_test_table(document: Document, items: Sequence[TestItem], settings: Settings) -> None:
    table = document.add_table(rows=1, cols=4)
    table.style = document.styles["Normal Table"]
    table.autofit = False
    _set_table_properties(table)

    header = table.rows[0]
    _set_row_height(header, settings.row_height_twips)
    for index, title in enumerate(settings.column_titles):
        _write_cell(
            header.cells[index],
            settings,
            text=title,
            bold=True,
            top_border=False,
            bottom_border=True,
            width=COLUMN_WIDTHS[index],
        )

    for item in items:
        row = table.add_row()
        _set_row_height(row, settings.row_height_twips)
        _write_cell(
            row.cells[0], settings, text=str(item.number), bold=True,
            align_right=True, width=COLUMN_WIDTHS[0],
        )
        _write_cell(row.cells[1], settings, text=item.german, width=COLUMN_WIDTHS[1])
        _write_cell(row.cells[2], settings, text=item.english, width=COLUMN_WIDTHS[2])
        _write_cell(
            row.cells[3],
            settings,
            text=item.sentence,
            highlight=item.sentence_form if settings.bold_target_word else "",
            width=COLUMN_WIDTHS[3],
        )


def build_document(pair: TestPair, settings: Settings | None = None) -> Document:
    """Baut das vollständige Word-Dokument mit beiden Tests."""
    settings = settings or Settings()
    document = Document()
    _configure_section(document)

    normal = document.styles["Normal"]
    normal.font.name = settings.font_name
    normal.font.size = Pt(settings.font_size_pt)

    _add_heading(document, settings.heading_test1)
    _add_test_table(document, pair.test1, settings)

    _add_heading(document, settings.heading_test2, settings.page_break_between_tests)
    _add_test_table(document, pair.test2, settings)

    return document


def _table_rows(items, settings: Settings) -> list[tuple[str, str, str, str]]:
    rows = [tuple(settings.column_titles)]
    rows += [(str(i.number), i.german, i.english, i.sentence) for i in items]
    return rows


def check_page_fit(items, settings: Settings | None = None) -> FitResult:
    """Passt eine Liste auf eine A4-Seite?"""
    settings = settings or Settings()
    return check_fits_page(
        _table_rows(items, settings),
        column_widths=COLUMN_WIDTHS,
        font_size_pt=settings.font_size_pt,
        row_height_twips=settings.row_height_twips,
        page_height=PAGE_HEIGHT,
        margin_top=MARGIN_TOP,
        margin_bottom=MARGIN_BOTTOM,
    )


def write_docx(
    pair: TestPair, target: str | Path | IO[bytes], settings: Settings | None = None
) -> None:
    """Schreibt die Word-Datei an den angegebenen Ort."""
    build_document(pair, settings).save(target)


def to_bytes(pair: TestPair, settings: Settings | None = None) -> bytes:
    """Die Word-Datei als Bytes - für den Download-Button in der Weboberfläche."""
    buffer = BytesIO()
    write_docx(pair, buffer, settings)
    return buffer.getvalue()


def suggested_filename(pair: TestPair) -> str:
    label = pair.unit_label or (
        "Starter_Unit" if pair.unit == 0 else f"Unit_{pair.unit}"
    )
    safe = re.sub(r"[^A-Za-z0-9_-]+", "_", label).strip("_") or "Unit"
    return f"Vocabulary_{safe}.docx"
